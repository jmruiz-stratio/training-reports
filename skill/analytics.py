from pathlib import Path
from datetime import date
import duckdb


def partner_kpis(workdir: str) -> list[dict]:
    users = Path(workdir) / "users.parquet"
    if not users.exists():
        return []
    q = f"""
    SELECT
      CASE WHEN lower(email) LIKE '%stratio%' OR lower(username) LIKE '%stratio%'
           THEN 'stratio' ELSE list_last(string_split(username, '-')) END AS partner,
      count(*) AS users
    FROM read_parquet('{users}')
    GROUP BY 1 ORDER BY 2 DESC
    """
    rows = duckdb.sql(q).fetchall()
    return [{"partner": p, "users": n} for p, n in rows]


def export_partner_kpis_csv(workdir: str, output_csv: str) -> str:
    rows = partner_kpis(workdir)
    import csv
    with open(output_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["partner", "users"])
        w.writeheader()
        w.writerows(rows)
    return output_csv


def _fp(workdir_path: Path, name: str) -> str | None:
    for f in workdir_path.glob(f"{name}*.parquet"):
        return str(f)
    return None


def _load_empleados_to_duckdb(con: duckdb.DuckDBPyConnection, xlsx_path: str) -> int:
    """Carga el Excel de empleados en raw_empleados. Devuelve nº de filas cargadas."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws_name = "Hoja2" if "Hoja2" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[ws_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return 0
    con.execute("""
        CREATE TABLE raw_empleados (
            id_vp VARCHAR, nombre VARCHAR, apellidos VARCHAR,
            lugar_trabajo VARCHAR, email VARCHAR,
            departamento VARCHAR, manager VARCHAR, rol VARCHAR
        )
    """)
    data = [
        (
            str(r[0] or ""), str(r[1] or ""), str(r[2] or ""), str(r[3] or ""),
            str(r[4] or "").lower().strip(),
            str(r[5] or ""), str(r[6] or ""), str(r[7] or ""),
        )
        for r in rows[1:]   # omitir cabecera
        if r[4]             # requerir email no vacío
    ]
    con.executemany("INSERT INTO raw_empleados VALUES (?,?,?,?,?,?,?,?)", data)
    return len(data)


def internal_course_compliance(
    workdir: str,
    course_pattern: str,
    output_excel: str = "reports/compliance.xlsx",
    partner: str = "stratio",
    empleados_xlsx: str | None = None,
) -> dict:
    """
    Informe de cumplimiento para un curso interno.

    Sin empleados_xlsx: universo = usuarios Moodle inscritos del partner.
    Con empleados_xlsx: universo = todos los empleados del fichero.
      Estado por empleado: completado / pendiente (inscrito sin completar) / sin_inscribir.
      Hojas extra: sin_inscribir, por_rol, por_departamento.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("Requiere openpyxl instalado (pip install openpyxl)")

    base = Path(workdir)
    if not base.exists():
        raise RuntimeError(f"Workdir no encontrado: {workdir}")

    users_p       = _fp(base, "users")
    courses_p     = _fp(base, "courses")
    completions_p = _fp(base, "course_completions")
    enrollments_p = _fp(base, "user_courses")
    grades_p      = _fp(base, "user_grades")

    if not users_p or not courses_p:
        raise RuntimeError("Faltan parquets obligatorios: users y courses")
    if not completions_p and not enrollments_p:
        raise RuntimeError("Se necesita course_completions o user_courses para el informe de cumplimiento")

    con = duckdb.connect()
    con.execute(f"CREATE VIEW raw_users   AS SELECT * FROM read_parquet('{users_p}')")
    con.execute(f"CREATE VIEW raw_courses AS SELECT * FROM read_parquet('{courses_p}')")
    if completions_p:
        con.execute(f"CREATE VIEW raw_cc AS SELECT * FROM read_parquet('{completions_p}')")
    if enrollments_p:
        con.execute(f"CREATE VIEW raw_uc AS SELECT * FROM read_parquet('{enrollments_p}')")
    if grades_p:
        con.execute(f"CREATE VIEW raw_grades AS SELECT * FROM read_parquet('{grades_p}')")

    safe_pattern = course_pattern.replace("'", "''").lower()
    safe_partner  = partner.replace("'", "''").lower()

    # 1. Cursos que coinciden con el patrón
    matching = con.execute(f"""
        SELECT id, fullname FROM raw_courses
        WHERE LOWER(fullname) LIKE '%{safe_pattern}%'
        ORDER BY fullname
    """).fetchall()

    if not matching:
        con.close()
        return {"error": f"No se encontró ningún curso con patrón '{course_pattern}'", "cursos": []}

    ids = ", ".join(str(c[0]) for c in matching)

    # Filtro de partner para Moodle (identifica usuarios del partner por email/username)
    if safe_partner == "stratio":
        moodle_partner_filter = "LOWER(u.email) LIKE '%stratio%' OR LOWER(u.username) LIKE '%stratio%'"
    else:
        moodle_partner_filter = f"list_last(string_split(u.username, '-')) = '{safe_partner}'"

    # 2. Tablas base: inscritos y completados en Moodle
    enrolled_parts = []
    if enrollments_p:
        enrolled_parts.append(f"SELECT userid, courseid FROM raw_uc WHERE courseid IN ({ids})")
    if completions_p:
        enrolled_parts.append(f"SELECT userid, courseid FROM raw_cc WHERE courseid IN ({ids})")
    enrolled_union = " UNION ".join(enrolled_parts)

    # t_enrolled: un email por usuario inscrito (normalizado)
    con.execute(f"""
        CREATE TABLE t_enrolled AS
        SELECT DISTINCT LOWER(u.email) AS email, u.id AS moodle_userid, u.username
        FROM ({enrolled_union}) e
        JOIN raw_users u ON u.id = e.userid
        WHERE {moodle_partner_filter}
    """)

    # t_completed: señal de completado.
    # Preferimos user_grades (graderaw IS NOT NULL) porque course_completions.completed
    # depende de que Moodle tenga el tracking de completado activado — muchos cursos no lo tienen.
    # Solo usamos course_completions como fallback si no hay notas para el curso.
    grades_for_course = (
        con.execute(f"SELECT COUNT(*) FROM raw_grades WHERE courseid IN ({ids})").fetchone()[0]
        if grades_p else 0
    )
    if grades_for_course > 0:
        con.execute(f"""
            CREATE TABLE t_completed AS
            SELECT
                LOWER(u.email) AS email,
                MIN(CASE WHEN g.gradedategraded > 0
                         THEN to_timestamp(g.gradedategraded) END) AS fecha_completado,
                MAX(CAST(g.graderaw AS DOUBLE))                     AS nota
            FROM raw_grades g
            JOIN raw_users u ON u.id = g.userid
            WHERE g.courseid IN ({ids}) AND g.graderaw IS NOT NULL
              AND ({moodle_partner_filter})
            GROUP BY LOWER(u.email)
        """)
    elif completions_p:
        con.execute(f"""
            CREATE TABLE t_completed AS
            SELECT
                LOWER(u.email) AS email,
                MIN(CASE WHEN cc.timecompleted > 0
                         THEN to_timestamp(cc.timecompleted) END) AS fecha_completado,
                NULL::DOUBLE                                        AS nota
            FROM raw_cc cc
            JOIN raw_users u ON u.id = cc.userid
            WHERE cc.courseid IN ({ids}) AND cc.completed = 1
              AND ({moodle_partner_filter})
            GROUP BY LOWER(u.email)
        """)
    else:
        con.execute("CREATE TABLE t_completed (email VARCHAR, fecha_completado TIMESTAMP, nota DOUBLE)")

    # ── Modo con fichero de empleados ─────────────────────────────────────────
    with_emp = bool(empleados_xlsx and Path(empleados_xlsx).exists())
    n_empleados = 0

    if with_emp:
        n_empleados = _load_empleados_to_duckdb(con, empleados_xlsx)

        # Estado de cada empleado
        con.execute("""
            CREATE TABLE emp_status AS
            SELECT
                emp.nombre,
                emp.apellidos,
                emp.email,
                emp.rol,
                emp.departamento,
                emp.manager,
                COALESCE(en.username, '') AS username_moodle,
                CASE
                    WHEN co.email IS NOT NULL THEN 'completado'
                    WHEN en.email IS NOT NULL THEN 'pendiente'
                    ELSE 'sin_inscribir'
                END AS estado,
                co.fecha_completado,
                co.nota
            FROM raw_empleados emp
            LEFT JOIN t_enrolled  en ON en.email = emp.email
            LEFT JOIN t_completed co ON co.email = emp.email
            ORDER BY estado, emp.apellidos, emp.nombre
        """)

        # Hojas de detalle
        completados_emp = con.execute("""
            SELECT nombre, apellidos, email, rol, departamento, manager,
                   username_moodle, nota, fecha_completado
            FROM emp_status WHERE estado = 'completado'
            ORDER BY apellidos, nombre
        """).fetchall()

        pendientes_emp = con.execute("""
            SELECT nombre, apellidos, email, rol, departamento, manager, username_moodle
            FROM emp_status WHERE estado = 'pendiente'
            ORDER BY apellidos, nombre
        """).fetchall()

        sin_inscribir_emp = con.execute("""
            SELECT nombre, apellidos, email, rol, departamento, manager
            FROM emp_status WHERE estado = 'sin_inscribir'
            ORDER BY apellidos, nombre
        """).fetchall()

        # Estadísticas por rol
        por_rol_rows = con.execute("""
            SELECT
                rol,
                COUNT(*)                                        AS total,
                COUNT(*) FILTER (WHERE estado = 'completado')  AS completados,
                COUNT(*) FILTER (WHERE estado = 'pendiente')   AS pendientes,
                COUNT(*) FILTER (WHERE estado = 'sin_inscribir') AS sin_inscribir,
                ROUND(100.0 * COUNT(*) FILTER (WHERE estado = 'completado')
                      / NULLIF(COUNT(*), 0), 1)                AS pct_completado
            FROM emp_status
            GROUP BY rol
            ORDER BY completados DESC, total DESC
        """).fetchall()

        # Estadísticas por departamento
        por_dept_rows = con.execute("""
            SELECT
                departamento,
                COUNT(*)                                        AS total,
                COUNT(*) FILTER (WHERE estado = 'completado')  AS completados,
                COUNT(*) FILTER (WHERE estado = 'pendiente')   AS pendientes,
                COUNT(*) FILTER (WHERE estado = 'sin_inscribir') AS sin_inscribir,
                ROUND(100.0 * COUNT(*) FILTER (WHERE estado = 'completado')
                      / NULLIF(COUNT(*), 0), 1)                AS pct_completado
            FROM emp_status
            GROUP BY departamento
            ORDER BY completados DESC, total DESC
        """).fetchall()

        # Resumen global
        resumen_rows = con.execute("""
            SELECT estado, COUNT(*) AS empleados,
                   ROUND(100.0 * COUNT(*) / SUM(COUNT(*)) OVER (), 1) AS pct_total
            FROM emp_status GROUP BY estado ORDER BY empleados DESC
        """).fetchall()

        n_comp = len(completados_emp)
        n_pend = len(pendientes_emp)
        n_si   = len(sin_inscribir_emp)
        con.close()

    else:
        # ── Modo solo Moodle (sin fichero de empleados) ───────────────────────
        if completions_p:
            completed_subq = f"""
            LEFT JOIN (
                SELECT userid, courseid,
                       CASE WHEN timecompleted > 0
                            THEN to_timestamp(timecompleted) END AS fecha_completado
                FROM raw_cc WHERE courseid IN ({ids}) AND completed = 1
            ) comp ON comp.userid = e.userid AND comp.courseid = e.courseid"""
            completed_sel = "comp.fecha_completado, (comp.userid IS NOT NULL) AS completado"
        else:
            completed_subq = ""
            completed_sel  = "NULL AS fecha_completado, FALSE AS completado"

        detail_sql = f"""
        WITH enrolled AS ({enrolled_union})
        SELECT
            u.id, u.username, u.email,
            u.firstname || ' ' || u.lastname AS nombre_completo,
            rc.fullname AS curso,
            {completed_sel},
            (u.suspended = 0) AS activo
        FROM enrolled e
        JOIN raw_users   u  ON u.id  = e.userid
        JOIN raw_courses rc ON rc.id = e.courseid
        {completed_subq}
        WHERE {moodle_partner_filter}
        ORDER BY completado DESC NULLS LAST, u.lastname, u.firstname
        """
        all_rows   = con.execute(detail_sql).fetchall()
        completados_moodle = [r for r in all_rows if r[6]]
        pendientes_moodle  = [r for r in all_rows if not r[6]]

        resumen_sql = f"""
        WITH enrolled AS ({enrolled_union})
        SELECT
            rc.fullname AS curso,
            COUNT(DISTINCT e.userid) AS total_inscritos,
            COUNT(DISTINCT CASE WHEN comp.userid IS NOT NULL THEN e.userid END) AS completados,
            COUNT(DISTINCT CASE WHEN comp.userid IS NULL     THEN e.userid END) AS pendientes,
            ROUND(100.0 * COUNT(DISTINCT CASE WHEN comp.userid IS NOT NULL THEN e.userid END)
                  / NULLIF(COUNT(DISTINCT e.userid), 0), 1) AS pct_completado
        FROM enrolled e
        JOIN raw_users   u  ON u.id  = e.userid
        JOIN raw_courses rc ON rc.id = e.courseid
        {"LEFT JOIN (SELECT userid, courseid FROM raw_cc WHERE courseid IN (" + ids + ") AND completed = 1) comp ON comp.userid = e.userid AND comp.courseid = e.courseid" if completions_p else ""}
        WHERE {moodle_partner_filter}
        GROUP BY rc.fullname ORDER BY rc.fullname
        """ if completions_p else f"""
        WITH enrolled AS ({enrolled_union})
        SELECT rc.fullname, COUNT(DISTINCT e.userid), 0, COUNT(DISTINCT e.userid), 0.0
        FROM enrolled e JOIN raw_users u ON u.id = e.userid
        JOIN raw_courses rc ON rc.id = e.courseid
        WHERE {moodle_partner_filter} GROUP BY rc.fullname
        """
        resumen_rows = con.execute(resumen_sql).fetchall()
        n_comp = len(completados_moodle)
        n_pend = len(pendientes_moodle)
        n_si   = 0
        con.close()

    # ── Excel ─────────────────────────────────────────────────────────────────
    HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
    HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
    ALT_FILL    = PatternFill("solid", fgColor="EBF3FB")
    NORMAL_FONT = Font(size=10)

    def _strip_tz(v):
        return v.replace(tzinfo=None) if hasattr(v, "tzinfo") and v.tzinfo else v

    def write_sheet(ws, rows, cols):
        ws.append(cols)
        for cell in ws[1]:
            cell.font      = HEADER_FONT
            cell.fill      = HEADER_FILL
            cell.alignment = Alignment(horizontal="center")
        for i, row in enumerate(rows, start=2):
            ws.append([_strip_tz(v) for v in row])
            if i % 2 == 0:
                for cell in ws[i]:
                    cell.fill = ALT_FILL
                    cell.font = NORMAL_FONT
        for col_idx in range(1, len(cols) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = max(
                (len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(1, ws.max_row + 1)),
                default=8,
            )
            ws.column_dimensions[col_letter].width = min(max_len + 2, 40)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    generated_sheets = []

    if with_emp:
        # Resumen global
        write_sheet(
            wb.create_sheet(title="resumen"),
            resumen_rows,
            ["estado", "empleados", "pct_total"],
        )
        generated_sheets.append("resumen")

        # Completados (con datos de RRHH)
        write_sheet(
            wb.create_sheet(title="completados"),
            completados_emp,
            ["nombre", "apellidos", "email", "rol", "departamento", "manager",
             "username_moodle", "nota", "fecha_completado"],
        )
        generated_sheets.append("completados")

        # Pendientes (inscritos pero sin completar)
        write_sheet(
            wb.create_sheet(title="pendientes"),
            pendientes_emp,
            ["nombre", "apellidos", "email", "rol", "departamento", "manager", "username_moodle"],
        )
        generated_sheets.append("pendientes")

        # Sin inscribir
        write_sheet(
            wb.create_sheet(title="sin_inscribir"),
            sin_inscribir_emp,
            ["nombre", "apellidos", "email", "rol", "departamento", "manager"],
        )
        generated_sheets.append("sin_inscribir")

        # Por rol
        write_sheet(
            wb.create_sheet(title="por_rol"),
            por_rol_rows,
            ["rol", "total", "completados", "pendientes", "sin_inscribir", "pct_completado"],
        )
        generated_sheets.append("por_rol")

        # Por departamento
        write_sheet(
            wb.create_sheet(title="por_departamento"),
            por_dept_rows,
            ["departamento", "total", "completados", "pendientes", "sin_inscribir", "pct_completado"],
        )
        generated_sheets.append("por_departamento")

    else:
        write_sheet(
            wb.create_sheet(title="resumen"),
            resumen_rows,
            ["curso", "total_inscritos", "completados", "pendientes", "pct_completado"],
        )
        generated_sheets.append("resumen")

        write_sheet(
            wb.create_sheet(title="completados"),
            [(r[1], r[2], r[3], r[4], r[5]) for r in completados_moodle],
            ["username", "email", "nombre_completo", "curso", "fecha_completado"],
        )
        generated_sheets.append("completados")

        write_sheet(
            wb.create_sheet(title="pendientes"),
            [(r[1], r[2], r[3], r[4]) for r in pendientes_moodle],
            ["username", "email", "nombre_completo", "curso"],
        )
        generated_sheets.append("pendientes")

    out_path = Path(output_excel)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)

    result = {
        "excel":              str(out_path.resolve()),
        "cursos_encontrados": [c[1] for c in matching],
        "partner":            partner,
        "fecha":              date.today().isoformat(),
        "completados":        n_comp,
        "pendientes":         n_pend,
        "sheets":             generated_sheets,
    }
    if with_emp:
        result["total_empleados"]  = n_empleados
        result["sin_inscribir"]    = n_si
        result["tasa_completado"]  = round(n_comp / n_empleados * 100, 1) if n_empleados else 0.0
    else:
        result["total_inscritos"]  = n_comp + n_pend
        result["tasa_completado"]  = round(n_comp / (n_comp + n_pend) * 100, 1) if (n_comp + n_pend) else 0.0
    return result


def _load_reporting_helpers():
    try:
        from src.local_validate import build_semantic, export_excel, reporting_queries
    except SystemExit as e:
        raise RuntimeError("Reporting requiere openpyxl instalado (pip install openpyxl)") from e
    return build_semantic, export_excel, reporting_queries


def export_reporting_package(
    workdir: str,
    output_excel: str,
    output_csv_dir: str,
    partner: str | None = None,
    practices_csv: str | None = None,
) -> dict:
    build_semantic, export_excel, reporting_queries = _load_reporting_helpers()

    base = Path(workdir)
    if not base.exists():
        raise RuntimeError(f"Workdir no encontrado: {workdir}")

    con = duckdb.connect()
    semantic_tables = build_semantic(con, base)

    if partner:
        p = partner.lower()
        for t in ["f_usuarios", "f_inscripciones", "f_certificaciones", "f_actividad", "f_dedicacion"]:
            if t in semantic_tables:
                con.execute(f"ALTER TABLE {t} RENAME TO {t}_all")
                con.execute(f"CREATE VIEW {t} AS SELECT * FROM {t}_all WHERE partner = '{p}'")

    output_excel_path = Path(output_excel)
    pcsv = Path(practices_csv) if practices_csv else Path("reports/sesiones_practicas.csv")
    sheets = export_excel(con, semantic_tables, output_excel_path, pcsv if pcsv.exists() else None)

    csv_dir = Path(output_csv_dir)
    csv_dir.mkdir(parents=True, exist_ok=True)
    exported_csv: list[str] = []

    for sheet_name, table_name, sql in reporting_queries(semantic_tables, con):
        con.execute(f"CREATE OR REPLACE TABLE {table_name} AS {sql}")
        out_csv = csv_dir / f"{sheet_name}.csv"
        con.execute(f"COPY (SELECT * FROM {table_name}) TO '{out_csv}' (HEADER, DELIMITER ',')")
        exported_csv.append(str(out_csv))

    con.close()

    return {
        "excel": str(output_excel_path),
        "sheets": sheets,
        "csv_files": exported_csv,
        "partner_filter": partner,
    }
